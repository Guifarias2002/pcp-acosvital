# Acrescenta o PRODUÇÃO SETEMBRO (só Pedido de Venda) ao CONTROLE GERAL do dia 19,
# mantendo o modelo (aba, colunas, formatação, fórmula) e gera o PDF no mesmo
# layout da "Faturamento - Apresentacao Diretoria".
import collections, copy, openpyxl
from classif import *

PASTA = r'C:/Users/guilherme.santos/Desktop/LEVANTAMENTO PRODUÇÃO CALDEIRARIA/'
MODELO = PASTA + 'CONTROLE GERAL DE FATURAMENTO AÇOS VITAL.xlsx'
NOVO_ARQ = r'C:/Users/guilherme.santos/Downloads/PRODUÇÃO SETEMBRO.xlsx'
SAIDA_X = PASTA + 'CONTROLE GERAL DE FATURAMENTO AÇOS VITAL - ATUALIZADO.xlsx'
SAIDA_P = PASTA + 'Faturamento - Apresentacao Diretoria - ATUALIZADO.pdf'

# grafia original (com os espaços do modelo) de cada valor normalizado
graf = [collections.defaultdict(collections.Counter) for _ in range(3)]
for r in ref:
    for i in range(3):
        if r[4 + i]: graf[i][norm(r[4 + i])][r[4 + i]] += 1
orig = lambda i, v: graf[i][v].most_common(1)[0][0] if graf[i].get(v) else v

nv = [r for r in openpyxl.load_workbook(NOVO_ARQ, data_only=True).worksheets[0].iter_rows(min_row=4, values_only=True) if r[0]]
# colunas do novo: Descrição | Total de Mercadoria | Total da Nota Fiscal | Pedido | Nota Fiscal
novas = sorted([r for r in nv if tipo_doc(r[3]) == 'Pedido de Venda'], key=lambda r: str(r[0]))
fora = [r for r in nv if tipo_doc(r[3]) != 'Pedido de Venda']

wb = openpyxl.load_workbook(MODELO)
ws = wb.worksheets[0]
ult = ws.max_row
while not ws.cell(ult, 1).value: ult -= 1
modelo_linha = ult - 1  # linha comum (a última é a Caldeiraria Pesada, em verde)
estilos = [copy.copy(ws.cell(modelo_linha, c)._style) for c in range(1, 8)]
linha = ult
classes = []
for r in novas:
    (f, l, e), como = classificar(r[0])
    classes.append((r, f, l, e, como))
    linha += 1
    vals = [str(r[0]), float(r[2] or 0), r[3], r[4], orig(0, f), orig(1, l), orig(2, e)]
    for c, v in enumerate(vals, 1):
        cel = ws.cell(linha, c, v)
        cel._style = copy.copy(estilos[c - 1])
ws['B3'] = f'=SUM(B4:B{max(2687, linha + 200)})'
ws.auto_filter.ref = f'A3:G{linha}'
wb.save(SAIDA_X)

# ── totais (todas as linhas: dia 19 + novas) ─────────────────────────────────
todas = [(float(r[1] or 0), norm(r[4]) or '(SEM FAMÍLIA)', norm(r[5]), norm(r[6])) for r in ref]
todas += [(float(r[2] or 0), f or '(SEM FAMÍLIA)', l, e) for r, f, l, e, _ in classes]
tot = sum(t[0] for t in todas)
E = collections.Counter(); L = collections.Counter(); F = collections.Counter(); LF = collections.defaultdict(collections.Counter)
for v, f, l, e in todas:
    E[e] += v; L[l] += v; F[f] += v; LF[l][f] += v
fab = tot - L['REVENDA']
v19 = sum(float(r[1] or 0) for r in ref); vnovo = sum(float(r[2] or 0) for r in novas)

# conferência
import openpyxl as _o
chk = [r for r in _o.load_workbook(SAIDA_X, data_only=False).worksheets[0].iter_rows(min_row=4, values_only=True) if r[0]]
assert abs(sum(float(r[1]) for r in chk) - tot) < 0.01
assert len(chk) == len(ref) + len(novas)

# ── PDF igual ao modelo ──────────────────────────────────────────────────────
import matplotlib; matplotlib.use('Agg')
import matplotlib.pyplot as plt
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import cm

brl = lambda v: ('R$ ' + f'{v:,.2f}').replace(',', 'X').replace('.', ',').replace('X', '.')
pct = lambda v: f'{v * 100:.1f}%'.replace('.', ',')
tit = lambda s: s.title().replace('Ç', 'ç').replace('(Sem Família)', '(Sem Família)')
AZ = colors.HexColor('#1A3A5C'); VD = colors.HexColor('#16794B'); LAR = colors.HexColor('#C2410C')
st = getSampleStyleSheet()
h1 = ParagraphStyle('h1', parent=st['Title'], alignment=0, textColor=AZ, fontSize=22, spaceAfter=2)
h2 = ParagraphStyle('h2', parent=st['Heading2'], textColor=AZ, fontSize=14, spaceBefore=10)
sub = ParagraphStyle('sub', parent=st['Normal'], textColor=colors.HexColor('#64748B'), fontSize=9.5)
small = ParagraphStyle('sm', parent=st['Normal'], textColor=colors.HexColor('#64748B'), fontSize=8.5, leading=11)
body = ParagraphStyle('b', parent=st['Normal'], fontSize=9.5, leading=13)
kpi = ParagraphStyle('k', parent=st['Normal'], fontSize=9.5, leading=21)

def barras(cont, fn, cores):
    itens = [(tit(k), v) for k, v in cont.most_common() if v > 0]
    fig, ax = plt.subplots(figsize=(8, 0.62 * len(itens) + 0.9))
    cs = cores if isinstance(cores, list) else [cores] * len(itens)
    ax.barh([i[0] for i in itens][::-1], [i[1] for i in itens][::-1], color=cs[:len(itens)][::-1])
    mx = max(i[1] for i in itens)
    for y, (_, v) in enumerate(itens[::-1]):
        ax.text(v + mx * 0.01, y, brl(v), va='center', fontsize=8, color='#1A3A5C', fontweight='bold')
    ax.set_xlim(0, mx * 1.3); ax.spines[['top', 'right']].set_visible(False)
    ax.xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda x, _: f'{x / 1e6:.0f}M' if mx >= 5e6 else f'{x / 1e6:.1f}M'.replace('.', ',')))
    plt.tight_layout(); fig.savefig(fn, dpi=150); plt.close(fig)
    return Image(fn, width=17 * cm, height=17 * cm * fig.get_figheight() / fig.get_figwidth())

def tabela(linhas_t, col_w, destaque=(), ultimo_escuro=True):
    t = Table(linhas_t, colWidths=col_w)
    sty = [('BACKGROUND', (0, 0), (-1, 0), AZ), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
           ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('FONTSIZE', (0, 0), (-1, -1), 9.5),
           ('ALIGN', (1, 0), (-1, -1), 'RIGHT'), ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F1F5F9')]),
           ('TOPPADDING', (0, 0), (-1, -1), 6), ('BOTTOMPADDING', (0, 0), (-1, -1), 6)]
    if ultimo_escuro:
        sty += [('BACKGROUND', (0, -1), (-1, -1), AZ), ('TEXTCOLOR', (0, -1), (-1, -1), colors.white), ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold')]
    for i, cor, bg in destaque:
        sty += [('TEXTCOLOR', (0, i), (-1, i), cor), ('FONTNAME', (0, i), (-1, i), 'Helvetica-Bold'), ('BACKGROUND', (0, i), (-1, i), bg)]
    t.setStyle(TableStyle(sty)); return t

doc = SimpleDocTemplate(SAIDA_P, pagesize=A4, leftMargin=1.8 * cm, rightMargin=1.8 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
S = [Paragraph('Faturamento — Apresentação à Diretoria', h1),
     Paragraph('Por Empresa de Fabricação e Local de Produção · valores por Total da Nota Fiscal', sub), Spacer(1, 10)]
kp = Table([[Paragraph(f'<font size=8 color="#64748B">FATURAMENTO TOTAL</font><br/><font size=16 color="#1A3A5C"><b>{brl(tot)}</b></font><br/><font size=7 color="#64748B">base: Total da Nota Fiscal</font>', kpi),
             Paragraph(f'<font size=8 color="#64748B">FABRICADO (produção própria)</font><br/><font size=16 color="#16794B"><b>{brl(fab)}</b></font><br/><font size=7 color="#64748B">{pct(fab / tot)} do total</font>', kpi),
             Paragraph(f'<font size=8 color="#64748B">REVENDA</font><br/><font size=16 color="#14532D"><b>{brl(L["REVENDA"])}</b></font><br/><font size=7 color="#64748B">{pct(L["REVENDA"] / tot)} do total</font>', kpi)]],
           colWidths=[5.8 * cm] * 3)
kp.setStyle(TableStyle([('BACKGROUND', (0, 0), (0, 0), colors.HexColor('#EEF4FB')), ('BACKGROUND', (1, 0), (1, 0), colors.HexColor('#EAF6EF')), ('BACKGROUND', (2, 0), (2, 0), colors.HexColor('#FDF1E8')),
                        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')), ('TOPPADDING', (0, 0), (-1, -1), 8), ('BOTTOMPADDING', (0, 0), (-1, -1), 8)]))
S += [kp, Paragraph('Por Empresa de Fabricação', h2), barras(E, 'e.png', ['#1A3A5C', '#1D4ED8', '#1D4ED8', '#1D4ED8'])]
S += [tabela([['Empresa de Fabricação', 'Valor (NF)', '% do Total']] + [[tit(k), brl(v), pct(v / tot)] for k, v in E.most_common()] + [['TOTAL GERAL', brl(tot), '100,0%']], [9 * cm, 5 * cm, 3 * cm])]

S += [PageBreak(), Paragraph('Por Local de Produção (fábricas)', h2)]
Lfab = collections.Counter({k: v for k, v in L.items() if k != 'REVENDA'})
azuis = ['#1A3A5C', '#1D4ED8', '#2563EB', '#3B82F6', '#60A5FA', '#93C5FD', '#BFDBFE', '#DBEAFE']
S += [barras(Lfab, 'l.png', azuis)]
rows = [['Local de Produção', 'Valor (NF)', '% do Fabricado']] + [[tit(k), brl(v), pct(v / fab)] for k, v in Lfab.most_common()]
rows += [['TOTAL FABRICADO', brl(fab), '100,0%'], ['Revenda (não fabricado)', brl(L['REVENDA']), '—'], ['TOTAL GERAL FATURADO', brl(tot), '—']]
S += [tabela(rows, [9 * cm, 5 * cm, 3 * cm], destaque=[(len(rows) - 3, VD, colors.HexColor('#EAF6EF')), (len(rows) - 2, LAR, colors.HexColor('#FDF1E8'))])]

S += [PageBreak(), Paragraph('Por Família (classificação do produto)', h2)]
S += [tabela([['Família', 'Valor (NF)', '% do Total']] + [[tit(k), brl(v), pct(v / tot)] for k, v in F.most_common()] + [['TOTAL', brl(tot), '100,0%']], [9 * cm, 5 * cm, 3 * cm])]

S += [PageBreak(), Paragraph('Detalhe — Família dentro de cada Fábrica', h2)]
for loc, v in Lfab.most_common():
    cab = Table([[Paragraph(f'<b><font color="#1A3A5C">{tit(loc)}</font></b>', body), Paragraph(f'<para alignment="right"><b><font color="#16794B">{brl(v)}</font></b></para>', body)]], colWidths=[11 * cm, 6 * cm])
    S += [cab, tabela([['Família', 'Valor (NF)']] + [[tit(k), brl(x)] for k, x in LF[loc].most_common()], [12 * cm, 5 * cm], ultimo_escuro=False), Spacer(1, 4)]
S += [Spacer(1, 8), Paragraph(f'“Fabricado” = todos os Locais de Produção exceto Revenda. Valores por Total da Nota Fiscal. Fonte: CONTROLE GERAL DE FATURAMENTO AÇOS VITAL.xlsx (relatório de 19/09, {brl(v19)}) + PRODUÇÃO SETEMBRO.xlsx (Pedido de Venda, NF 52.408–52.799, {brl(vnovo)}).', small)]
doc.build(S)

tf = collections.Counter()
for r in fora: tf[tipo_doc(r[3])] += float(r[2] or 0)
print('dia19', round(v19, 2), 'novo PV', round(vnovo, 2), len(novas), 'linhas | total', round(tot, 2), 'fab', round(fab, 2))
print('fora do criterio no novo:', {k: round(v, 2) for k, v in tf.items()})
print('Locais:', {k: round(v, 2) for k, v in L.most_common()})
cc = collections.Counter(); cv = collections.Counter()
for r, f, l, e, como in classes: k = como.split(':')[0]; cc[k] += 1; cv[k] += float(r[2] or 0)
print({k: (cc[k], round(cv[k], 2)) for k in cc})
print(SAIDA_X); print(SAIDA_P)
