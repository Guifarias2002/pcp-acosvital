# Faturamento de SETEMBRO (só o export do mês, sem somar o relatório de agosto)
# no mesmo layout da apresentação, SEM a família Flanges (enviada à parte).
# Caldeiraria = planilha + SISTEMA (PCP Caldeiraria, mesmo valor da tela Valores
# por Mês: item não cancelado, mês da chegada na Caldeiraria). Pedido com valor no
# sistema que também está na planilha → fica o do sistema (linhas da planilha saem).
# Antes: node --env-file=.env.local scripts/faturamento/sistema_caldeiraria.mjs <CALD>
import collections, json, re, sys, openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from classif import classificar, tipo_doc
from pdf_modelo import gerar, PASTA

ARQ = r'C:/Users/guilherme.santos/Downloads/pivot.xlsx'
SAIDA_P = PASTA + 'Faturamento - Setembro (sem Flanges).pdf'
SAIDA_X = PASTA + 'FATURAMENTO SETEMBRO (sem Flanges) - CLASSIFICADO.xlsx'
CALD = sys.argv[1]
MES = '2026-09'
EMP = {'acosvital': 'AÇOS VITAL', 'hrm': 'HRM'}
num = lambda s: re.sub(r'\D', '', str(s or '')).lstrip('0')

nv = [r for r in openpyxl.load_workbook(ARQ, data_only=True).worksheets[0].iter_rows(min_row=4, values_only=True) if r[0]]
# colunas: Descrição | Total de Mercadoria | Total da Nota Fiscal | Pedido | Nota Fiscal
pv = [r for r in nv if tipo_doc(r[3]) == 'Pedido de Venda']
fora = [r for r in nv if tipo_doc(r[3]) != 'Pedido de Venda']
linhas = []
for r in pv:
    (f, l, e), como = classificar(r[0])
    linhas.append(dict(desc=str(r[0]), v=float(r[2] or 0), ped=r[3], nf=str(r[4]).strip(), fam=f or '(SEM FAMÍLIA)', loc=l, emp=e, como=como))
flanges = [x for x in linhas if x['fam'] == 'FLANGES']
sem = [x for x in linhas if x['fam'] != 'FLANGES']

# ── Caldeiraria do sistema ───────────────────────────────────────────────────
cald = [c for c in json.load(open(CALD, encoding='utf-8'))
        if c['status'] != 'cancelado' and (c['chegada'] or c['criado_em'])[:7] == MES]
sis = [dict(desc=f"PV {c['pedido']} - {c['cliente'] or ''} - {c['material']}", v=float(c['valor'] or 0), ped=f"Pedido de Venda nº {c['pedido']}",
            nf='', fam='CALDEIRARIA', loc='CALDEIRARIA (PLANEJAMENTO)', emp=EMP.get(c['empresa'], 'AÇOS VITAL'),
            como=f"sistema — {c['status']}" + (f", finalizado {c['finalizado_em']}" if c['finalizado_em'] else ''))
       for c in cald]
com_valor = {num(c['pedido']) for c in cald if (c['valor'] or 0) > 0}
dup = [x for x in sem if num(x['ped']) in com_valor]
sem = [x for x in sem if num(x['ped']) not in com_valor] + sis
vsis = sum(x['v'] for x in sis); vdup = sum(x['v'] for x in dup)

nfs = sorted(int(x['nf']) for x in linhas if x['nf'].isdigit())
fmt = lambda n: f'{n:,}'.replace(',', '.')
brl = lambda v: ('R$ ' + f'{v:,.2f}').replace(',', 'X').replace('.', ',').replace('X', '.')
vfl = sum(x['v'] for x in flanges)
tot, fab = gerar([(x['v'], x['fam'], x['loc'], x['emp'], x['nf']) for x in sem], SAIDA_P,
                 'Setembro · <b>sem Flanges</b> · planilha Omie + Caldeiraria do sistema · Por Empresa de Fabricação e Local de Produção · valores por Total da Nota Fiscal',
                 f'Fonte: export Omie de setembro (só Pedido de Venda, NFs {fmt(nfs[0])} a {fmt(nfs[-1])}), mesmos critérios do CONTROLE GERAL. '
                 f'Flanges ({brl(vfl)}) retirados — enviados à parte. Caldeiraria (Planejamento) = sistema PCP Caldeiraria ({brl(vsis)}, {len({num(c["pedido"]) for c in cald})} pedidos, mês de chegada na Caldeiraria); '
                 f'{brl(vdup)} da planilha de pedidos que já estão no sistema foram retirados para não contar duas vezes.')

# ── planilha p/ conferência ──────────────────────────────────────────────────
wb = openpyxl.Workbook()
H = Font(bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1A3A5C'); amarelo = PatternFill('solid', fgColor='FEF3C7')
cab = ['Descrição do Produto (completa)', 'Total da Nota Fiscal', 'Pedido', 'Nota Fiscal', 'Familia', 'Local de produção', 'Empresa de fabricação', 'Como foi classificado']

def aba(ws, rows):
    ws.append(cab)
    for c in ws[1]: c.font = H; c.fill = HF
    for x in sorted(rows, key=lambda x: x['desc']):
        ws.append([x['desc'], x['v'], x['ped'], x['nf'], x['fam'], x['loc'], x['emp'], x['como']])
        ws.cell(ws.max_row, 2).number_format = '#,##0.00'
        if 'CONFERIR' in x['como']:
            for c in ws[ws.max_row]: c.fill = amarelo
    ws.append(['TOTAL', sum(x['v'] for x in rows)]); ws.cell(ws.max_row, 2).number_format = '#,##0.00'
    for c in ws[ws.max_row]: c.font = Font(bold=True)
    ws.auto_filter.ref = f'A1:H{ws.max_row - 1}'; ws.freeze_panes = 'A2'
    for i, w in enumerate([80, 18, 26, 12, 18, 28, 20, 44], 1): ws.column_dimensions[get_column_letter(i)].width = w

ws = wb.active; ws.title = 'Setembro sem Flanges'; aba(ws, sem)
aba(wb.create_sheet('Caldeiraria (sistema)'), sis)
aba(wb.create_sheet('Duplicados (retirados)'), dup)
aba(wb.create_sheet('Flanges (retirados)'), flanges)
s = wb.create_sheet('Fora do critério')
s.append(['Tipo de documento', 'Descrição do Produto (completa)', 'Total da Nota Fiscal', 'Pedido', 'Nota Fiscal'])
for c in s[1]: c.font = H; c.fill = HF
for r in sorted(fora, key=lambda r: (tipo_doc(r[3]), str(r[0]))):
    s.append([tipo_doc(r[3]), str(r[0]), float(r[2] or 0), r[3], r[4]]); s.cell(s.max_row, 3).number_format = '#,##0.00'
for col, w in zip('ABCDE', [28, 80, 18, 30, 12]): s.column_dimensions[col].width = w
wb.save(SAIDA_X)

assert abs(tot + vfl + vdup - vsis - sum(x['v'] for x in linhas)) < 0.01
print(SAIDA_P); print(SAIDA_X)
print('total sem flanges', round(tot, 2), '| fabricado', round(fab, 2), '| flanges', round(vfl, 2), '| sistema', round(vsis, 2), '| duplicados', round(vdup, 2))
for x in dup: print('  dup', x['ped'], x['nf'], round(x['v'], 2), x['desc'][:70])
