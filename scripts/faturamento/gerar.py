import collections, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from classif import *

OUT = r'C:/Users/guilherme.santos/Desktop/LEVANTAMENTO PRODUÇÃO CALDEIRARIA/'
MES = 'Setembro 2026'

novo = [r for r in openpyxl.load_workbook(NOVO, data_only=True).worksheets[0].iter_rows(min_row=4, values_only=True) if r[0]]
pv = [r for r in novo if tipo_doc(r[2]) == 'Pedido de Venda']
fora = [r for r in novo if tipo_doc(r[2]) != 'Pedido de Venda']
linhas = []
for r in pv:
    (f, l, e), m = classificar(r[0])
    linhas.append(dict(desc=str(r[0]), v=float(r[1] or 0), ped=r[2], nf=r[3], fam=f, loc=l, emp=e, como=m))

def agg(rows, key):
    c = collections.Counter()
    for x in rows: c[key(x)] += x['v']
    return c

tot = sum(x['v'] for x in linhas)
L = agg(linhas, lambda x: x['loc']); E = agg(linhas, lambda x: x['emp']); F = agg(linhas, lambda x: x['fam'])
fab = tot - L['REVENDA']

# anterior (mesmos critérios)
refL = collections.Counter(); refE = collections.Counter(); refF = collections.Counter()
for r in ref:
    v = float(r[1] or 0); refL[norm(r[5])] += v; refE[norm(r[6])] += v; refF[norm(r[4]) or '(SEM FAMÍLIA)'] += v
reftot = sum(refL.values()); reffab = reftot - refL['REVENDA']
ref_nf = sorted(int(r[3]) for r in ref if str(r[3]).isdigit() and int(r[3]) > 1000)

tit = lambda s: s.title().replace('Ç', 'ç').replace(' De ', ' de ').replace('Aços Vital', 'Aços Vital')

# ── Excel ───────────────────────────────────────────────────────────────────
wb = openpyxl.Workbook()
H = Font(bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1A3A5C')
money = '#,##0.00'
ws = wb.active; ws.title = 'Faturamento por Produto'
ws.append([f'Faturamento por Produto — {MES} (Pedido de Venda) — classificado com os critérios do CONTROLE GERAL'])
ws['A1'].font = Font(bold=True, size=13)
cab = ['Descrição do Produto (completa)', 'Total da Nota Fiscal', 'Pedido', 'Nota Fiscal', 'Familia ', 'Local de produção ', 'Empresa de fabricação', 'Como foi classificado']
ws.append(cab)
for c in ws[2]: c.font = H; c.fill = HF
ws.append([None, f'=SUM(B4:B{len(linhas) + 3})'])
ws['B3'].number_format = money; ws['B3'].font = Font(bold=True)
amarelo = PatternFill('solid', fgColor='FEF3C7')
for x in sorted(linhas, key=lambda x: x['desc']):
    ws.append([x['desc'], x['v'], x['ped'], x['nf'], x['fam'], x['loc'], x['emp'], x['como']])
    ws.cell(ws.max_row, 2).number_format = money
    if 'CONFERIR' in x['como'] or 'SEM REGRA' in x['como']:
        for c in ws[ws.max_row]: c.fill = amarelo
ws.auto_filter.ref = f'A2:H{ws.max_row}'
ws.freeze_panes = 'A3'
for i, w in enumerate([80, 18, 26, 12, 18, 28, 20, 44], 1): ws.column_dimensions[get_column_letter(i)].width = w

def aba_resumo(nome, cont, total, rot):
    s = wb.create_sheet(nome)
    s.append([rot, 'Valor (NF)', '% do total']); [setattr(c, 'font', H) or setattr(c, 'fill', HF) for c in s[1]]
    for k, v in cont.most_common():
        s.append([k or '(SEM FAMÍLIA)', v, v / total if total else 0])
        s.cell(s.max_row, 2).number_format = money; s.cell(s.max_row, 3).number_format = '0.0%'
    s.append(['TOTAL', total, 1]); s.cell(s.max_row, 2).number_format = money; s.cell(s.max_row, 3).number_format = '0.0%'
    for c in s[s.max_row]: c.font = Font(bold=True)
    s.column_dimensions['A'].width = 32; s.column_dimensions['B'].width = 20; s.column_dimensions['C'].width = 12

aba_resumo('Por Empresa', E, tot, 'Empresa de fabricação')
aba_resumo('Por Local de Produção', L, tot, 'Local de produção')
aba_resumo('Por Família', F, tot, 'Família')

s = wb.create_sheet('Comparativo x anterior')
s.append(['Local de produção', 'Anterior (NF 50.0xx–51.697)', '% anterior', MES, f'% {MES}'])
for c in s[1]: c.font = H; c.fill = HF
for k in sorted(set(refL) | set(L), key=lambda k: -(L[k] + refL[k])):
    s.append([k, refL[k], refL[k] / reftot, L[k], L[k] / tot])
    for col in (2, 4): s.cell(s.max_row, col).number_format = money
    for col in (3, 5): s.cell(s.max_row, col).number_format = '0.0%'
s.append(['TOTAL', reftot, 1, tot, 1])
for col in (2, 4): s.cell(s.max_row, col).number_format = money
for c in s[s.max_row]: c.font = Font(bold=True)
for i, w in enumerate([32, 24, 12, 22, 14], 1): s.column_dimensions[get_column_letter(i)].width = w

s = wb.create_sheet('Fora do critério')
s.append(['Tipo de documento', 'Descrição do Produto (completa)', 'Total da Nota Fiscal', 'Pedido', 'Nota Fiscal'])
for c in s[1]: c.font = H; c.fill = HF
for r in sorted(fora, key=lambda r: (tipo_doc(r[2]), str(r[0]))):
    s.append([tipo_doc(r[2]), str(r[0]), float(r[1] or 0), r[2], r[3]]); s.cell(s.max_row, 3).number_format = money
s.column_dimensions['A'].width = 28; s.column_dimensions['B'].width = 80; s.column_dimensions['C'].width = 18; s.column_dimensions['D'].width = 30
xlsx = OUT + f'FATURAMENTO {MES.upper()} - CLASSIFICADO.xlsx'
wb.save(xlsx)

# ── PDF (mesmo layout da apresentação anterior) ─────────────────────────────
import matplotlib; matplotlib.use('Agg')
import matplotlib.pyplot as plt
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import cm

brl = lambda v: ('R$ ' + f'{v:,.2f}').replace(',', 'X').replace('.', ',').replace('X', '.')
pct = lambda v: f'{v * 100:.1f}%'.replace('.', ',')
AZ = colors.HexColor('#1A3A5C'); VD = colors.HexColor('#16794B'); LAR = colors.HexColor('#C2410C')
st = getSampleStyleSheet()
h1 = ParagraphStyle('h1', parent=st['Title'], alignment=0, textColor=AZ, fontSize=22, spaceAfter=2)
h2 = ParagraphStyle('h2', parent=st['Heading2'], textColor=AZ, fontSize=14, spaceBefore=10)
sub = ParagraphStyle('sub', parent=st['Normal'], textColor=colors.HexColor('#64748B'), fontSize=9.5)
small = ParagraphStyle('sm', parent=st['Normal'], textColor=colors.HexColor('#64748B'), fontSize=8.5, leading=11)
body = ParagraphStyle('b', parent=st['Normal'], fontSize=9.5, leading=13)

def barras(cont, fn, cor='#1D4ED8'):
    itens = [(tit(k or '(Sem família)'), v) for k, v in cont.most_common() if v > 0]
    fig, ax = plt.subplots(figsize=(8, 0.55 * len(itens) + 0.8))
    ax.barh([i[0] for i in itens][::-1], [i[1] for i in itens][::-1], color=cor)
    mx = max(i[1] for i in itens)
    for y, (_, v) in enumerate(itens[::-1]):
        ax.text(v + mx * 0.01, y, brl(v), va='center', fontsize=8, color='#1A3A5C', fontweight='bold')
    ax.set_xlim(0, mx * 1.3); ax.spines[['top', 'right']].set_visible(False)
    ax.xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda x, _: (f'{x / 1e6:.0f}M' if mx >= 5e6 else f'{x / 1e3:.0f} mil')))
    plt.tight_layout(); fig.savefig(fn, dpi=150); plt.close(fig)
    return Image(fn, width=17 * cm, height=17 * cm * fig.get_figheight() / fig.get_figwidth())

def tabela(linhas_t, col_w, destaque=()):
    t = Table(linhas_t, colWidths=col_w)
    sty = [('BACKGROUND', (0, 0), (-1, 0), AZ), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
           ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('FONTSIZE', (0, 0), (-1, -1), 9),
           ('ALIGN', (1, 0), (-1, -1), 'RIGHT'), ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#F1F5F9')]),
           ('BACKGROUND', (0, -1), (-1, -1), AZ), ('TEXTCOLOR', (0, -1), (-1, -1), colors.white), ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
           ('TOPPADDING', (0, 0), (-1, -1), 5), ('BOTTOMPADDING', (0, 0), (-1, -1), 5)]
    for i, cor in destaque: sty += [('TEXTCOLOR', (0, i), (-1, i), cor), ('FONTNAME', (0, i), (-1, i), 'Helvetica-Bold')]
    t.setStyle(TableStyle(sty)); return t

pdf = OUT + f'Faturamento {MES} - Apresentacao Diretoria.pdf'
doc = SimpleDocTemplate(pdf, pagesize=A4, leftMargin=1.8 * cm, rightMargin=1.8 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
S = []
S += [Paragraph(f'Faturamento — {MES}', h1), Paragraph('Por Empresa de Fabricação e Local de Produção · valores por Total da Nota Fiscal · só Pedido de Venda (mesmo critério do relatório anterior)', sub), Spacer(1, 10)]
kp = Table([[Paragraph(f'<font size=8 color="#64748B">FATURAMENTO TOTAL</font><br/><font size=16 color="#1A3A5C"><b>{brl(tot)}</b></font><br/><font size=7 color="#64748B">{len(linhas)} linhas · Pedido de Venda</font>', body),
             Paragraph(f'<font size=8 color="#64748B">FABRICADO (produção própria)</font><br/><font size=16 color="#16794B"><b>{brl(fab)}</b></font><br/><font size=7 color="#64748B">{pct(fab / tot)} do total</font>', body),
             Paragraph(f'<font size=8 color="#64748B">REVENDA</font><br/><font size=16 color="#14532D"><b>{brl(L["REVENDA"])}</b></font><br/><font size=7 color="#64748B">{pct(L["REVENDA"] / tot)} do total</font>', body)]],
           colWidths=[5.8 * cm] * 3)
kp.setStyle(TableStyle([('BACKGROUND', (0, 0), (0, 0), colors.HexColor('#EEF4FB')), ('BACKGROUND', (1, 0), (1, 0), colors.HexColor('#EAF6EF')), ('BACKGROUND', (2, 0), (2, 0), colors.HexColor('#FDF1E8')),
                        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')), ('TOPPADDING', (0, 0), (-1, -1), 8), ('BOTTOMPADDING', (0, 0), (-1, -1), 8)]))
S += [kp, Paragraph('Por Empresa de Fabricação', h2), barras(E, 'e.png', '#1A3A5C')]
S += [tabela([['Empresa de Fabricação', 'Valor (NF)', '% do Total']] + [[tit(k), brl(v), pct(v / tot)] for k, v in E.most_common()] + [['TOTAL GERAL', brl(tot), '100,0%']], [9 * cm, 5 * cm, 3 * cm])]
S += [PageBreak(), Paragraph('Por Local de Produção (fábricas)', h2)]
Lfab = collections.Counter({k: v for k, v in L.items() if k != 'REVENDA'})
S += [barras(Lfab, 'l.png', '#1D4ED8')]
rows = [['Local de Produção', 'Valor (NF)', '% do Fabricado']] + [[tit(k), brl(v), pct(v / fab)] for k, v in Lfab.most_common()]
rows += [['TOTAL FABRICADO', brl(fab), '100,0%'], ['Revenda (não fabricado)', brl(L['REVENDA']), '—'], ['TOTAL GERAL FATURADO', brl(tot), '—']]
S += [tabela(rows, [9 * cm, 5 * cm, 3 * cm], destaque=[(len(rows) - 3, VD), (len(rows) - 2, LAR)])]
S += [PageBreak(), Paragraph('Por Família (classificação do produto)', h2)]
S += [tabela([['Família', 'Valor (NF)', '% do Total']] + [[tit(k or '(Sem família)'), brl(v), pct(v / tot)] for k, v in F.most_common()] + [['TOTAL', brl(tot), '100,0%']], [9 * cm, 5 * cm, 3 * cm])]
S += [PageBreak(), Paragraph('Detalhe — Família dentro de cada Fábrica', h2)]
for loc, v in Lfab.most_common():
    fams = agg([x for x in linhas if x['loc'] == loc], lambda x: x['fam'])
    S += [Paragraph(f'<b>{tit(loc)}</b> — <font color="#16794B"><b>{brl(v)}</b></font>', body)]
    S += [tabela([['Família', 'Valor (NF)']] + [[tit(k or '(Sem família)'), brl(x)] for k, x in fams.most_common()] + [['Total', brl(v)]], [12 * cm, 5 * cm]), Spacer(1, 6)]

# Comparativo + o que mudou
S += [PageBreak(), Paragraph(f'{MES} × relatório anterior', h2),
      Paragraph(f'Anterior = CONTROLE GERAL DE FATURAMENTO (NFs {ref_nf[0]:,}–{ref_nf[-1]:,}). Setembro começa depois — <b>nenhuma nota repetida</b>. Mesmos critérios de Família / Local / Empresa.'.replace(',', '.'), small), Spacer(1, 6)]
rows = [['Local de Produção', 'Anterior', '%', MES, '%']]
for k in sorted(set(refL) | set(L), key=lambda k: -(L[k] + refL[k])):
    rows.append([tit(k), brl(refL[k]), pct(refL[k] / reftot), brl(L[k]), pct(L[k] / tot)])
rows.append(['TOTAL', brl(reftot), '100,0%', brl(tot), '100,0%'])
S += [tabela(rows, [5.6 * cm, 3.7 * cm, 1.8 * cm, 3.7 * cm, 1.8 * cm])]
tfora = collections.Counter()
for r in fora: tfora[tipo_doc(r[2])] += float(r[1] or 0)
n_cod = sum(1 for x in linhas if x['como'].startswith('código')); v_cod = sum(x['v'] for x in linhas if x['como'].startswith('código'))
S += [Spacer(1, 10), Paragraph('<b>O que tem de novo neste mês</b>', body), Spacer(1, 3)]
itens = [
    f'O export de setembro traz outros tipos de documento que o relatório anterior não tinha: ' + '; '.join(f'<b>{k}</b> {brl(v)}' for k, v in tfora.most_common()) + '. Pelo mesmo critério (só Pedido de Venda) ficaram <b>fora</b> dos totais — estão na aba "Fora do critério" da planilha.',
    f'Total do arquivo (tudo somado): {brl(tot + sum(tfora.values()))}. Faturamento por Pedido de Venda: <b>{brl(tot)}</b>.',
    f'{n_cod} linhas ({brl(v_cod)}) usam <b>exatamente a classificação do relatório anterior</b> (mesmo código de produto); as demais foram pelas mesmas regras (flange = Fábrica Flanges, grade/degrau = Fábrica de Grade, spool/suporte/estrutura/desenho = Caldeiraria Leve, telha sanduíche = Fábrica de Telha/Uberaba, chapa expandida = Terceiros, placa de desgaste 190/390 = Usinagem, resto = Revenda por família).',
    f'<b>Caldeiraria Pesada</b>: no relatório anterior entrou como uma linha lançada à mão (R$ 2.432.408,00), não vem do Omie. Em setembro <b>não está incluída</b> — falta o valor.',
    'Linhas em amarelo na planilha = conferir (chapa/placa "conforme desenho" — no anterior contou como Chapas/Revenda; e itens sem regra, lançados como Diversos).',
]
for t in itens: S += [Paragraph('• ' + t, body), Spacer(1, 3)]
S += [Spacer(1, 8), Paragraph('“Fabricado” = todos os Locais de Produção exceto Revenda. Valores por Total da Nota Fiscal. Fonte: export Omie de setembro (pivot) + critérios do CONTROLE GERAL DE FATURAMENTO AÇOS VITAL.xlsx.', small)]
doc.build(S)
print(xlsx); print(pdf); print(round(tot, 2), round(fab, 2))
