# Classifica o faturamento (export Omie) com os MESMOS critérios do
# "CONTROLE GERAL DE FATURAMENTO AÇOS VITAL.xlsx" (Família / Local de produção /
# Empresa de fabricação). 1) código já classificado no antigo; 2) regras por
# palavra (derivadas do antigo); 3) sobra = DIVERSOS/REVENDA (marcado p/ revisar).
import openpyxl, collections, re, json, sys, unicodedata

REF = r'C:/Users/guilherme.santos/Desktop/LEVANTAMENTO PRODUÇÃO CALDEIRARIA/CONTROLE GERAL DE FATURAMENTO AÇOS VITAL.xlsx'
NOVO = r'C:/Users/guilherme.santos/Downloads/pivot (2).xlsx'

def sem_acento(s):
    return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn')

def cod_de(desc):
    return str(desc).split(' - ')[0].strip().upper()

def tipo_doc(ped):
    return re.sub(r'\s*n\S*\s*\d+.*', '', str(ped or '')).strip()

ref = [r for r in openpyxl.load_workbook(REF, data_only=True).worksheets[0].iter_rows(min_row=4, values_only=True) if r[0]]
norm = lambda x: (x or '').strip()

# código → classificação (a de MAIOR valor quando o antigo teve mais de uma)
acc = collections.defaultdict(lambda: collections.Counter())
for r in ref:
    acc[cod_de(r[0])][(norm(r[4]), norm(r[5]), norm(r[6]))] += float(r[1] or 0) + 0.01
MAPA = {c: cnt.most_common(1)[0][0] for c, cnt in acc.items()}

REV = ('REVENDA', 'AÇOS VITAL')
def C(fam, local='REVENDA', emp='AÇOS VITAL'):
    return (fam, local, emp)

def regra(desc):
    d = sem_acento(str(desc).split(' - ', 1)[-1].upper())
    w = re.findall(r'[A-Z0-9]+', d)
    first = w[0] if w else ''
    has = lambda *k: any(re.search(r'\b' + x + r'\b', d) for x in k)
    plastico = has('PVC', 'CPVC', 'PVDF', 'PTFE', 'NYLON', 'POLIPROPILENO', 'PP', 'POLIACETAL', 'POLIETILENO', 'TEFLON', 'UHMW[0-9]*', 'ACRILICO')
    # Flange (inclusive "p/ tubo PEAD" e PRFV) = Fábrica Flanges — só flange de plástico vai p/ Plásticos (critério do relatório anterior)
    if first in ('FLANGE', 'FLANGES', 'DISCO', 'CONTRA', 'RAQUETE') and not plastico:
        return C('FLANGES', 'FABRICA FLANGES', 'HRM'), 'flange'
    if first == 'TUBO' and has('FLANGEADO') and has('PRFV'):
        return C('FLANGES', 'FABRICA FLANGES', 'HRM'), 'tubo flangeado PRFV'
    if first in ('GRADE', 'DEGRAU'): return C('GRADE DE PISO', 'FABRICA DE GRADE DE PISO', 'AÇOS VITAL'), 'grade/degrau'
    if first in ('VALVULA', 'REGISTRO', 'PURGADOR', 'VENTOSA'):
        return (C('PLASTICOS'), 'válvula plástica') if has('PVC', 'CPVC', 'PVDF') else (C('VALVULA'), 'válvula')
    # placa de desgaste 190/390 (Hardox/desgaste) = Fábrica Usinagem; chapa Hardox/Usiar inteira = revenda
    if first in ('CHAPA', 'PLACA') and (has('DESGASTE', 'ANTIDESGASTE', 'HARDOX') and re.search(r'(190|390)\s*MM', d)):
        return C('HARDOX', 'FABRICA USINAGEM', 'HRM'), 'placa de desgaste 190/390'
    if first in ('CHAPA', 'PLACA') and has('DESENHO', 'CONF', 'PROJETO') and not has('EXPANDIDA', 'PERFURADA'):
        return C('CHAPAS'), 'chapa conforme desenho (CONFERIR)'
    if first == 'KIT' and has('CHAPA'): return C('CHAPAS'), 'kit chapa'
    if first == 'KIT' and has('TUBULACAO', 'TUBO'): return C('TUBOS'), 'kit tubulação'
    # materiais que mandam na família
    if has('PRFV', 'FIBRA DE VIDRO'): return C('FIBRA DE VIDRO'), 'PRFV'
    if has('PEAD', 'PE100', 'PE 100', 'PE80') or first in ('COLARINHO',): return C('PEAD'), 'PEAD'
    if has('PVC', 'CPVC', 'PVDF', 'PTFE', 'NYLON', 'POLIPROPILENO', 'PP', 'POLIACETAL', 'POLIETILENO', 'TEFLON', 'UHMW[0-9]*', 'ACRILICO'): return C('PLASTICOS'), 'plástico'
    # caldeiraria leve (sob projeto/desenho)
    if first in ('SPOOL', 'SUPORTE', 'ESTRUTURA', 'BLANK', 'POCO', 'CAMISA', 'CAMISAS', 'BANCADA', 'SKID', 'TANQUE', 'CARRETEL', 'COIFA', 'DUTO', 'PLATAFORMA', 'GUARDA', 'ESCADA', 'PASSARELA', 'CORRIMAO') \
       or (has('CONFORME PROJETO', 'CONF PROJETO', 'CONFORME DESENHO', 'CONF DESENHO', 'CONF. PROJETO', 'CONF. DESENHO') and first not in ('FLANGE', 'ANEL', 'CALHA')):
        return C('CALDEIRARIA', 'CALDEIRARIA LEVE', 'HRM'), 'caldeiraria (projeto/desenho)'
    if first in ('FABRICACAO', 'TRELICA') or (first == 'CONJUNTO' and has('DESENHO')):
        return C('CALDEIRARIA', 'CALDEIRARIA LEVE', 'HRM'), 'caldeiraria (fabricação/desenho)'
    if first == 'RAQUETE': return C('FLANGES', 'FABRICA FLANGES', 'HRM'), 'raquete (flange)'
    if first in ('RED', 'NIPOLET', 'COLAR'): return C('CONEXÕES'), 'conexão'
    if first in ('PC', 'PC/PA'): return C('FIXAÇÃO'), 'fixação (estojo)'
    if first in ('VENTOSA',): return C('VALVULA'), 'válvula'
    if first == 'SUCATA': return C('DIVERSOS'), 'sucata'
    if first == 'TUBO' and has('CALANDRADO'): return C('TUBOS', 'CALDEIRARIA LEVE', 'HRM'), 'tubo calandrado'
    if first in ('FLANGE', 'FLANGES', 'DISCO', 'CONTRA'): return C('FLANGES', 'FABRICA FLANGES', 'HRM'), 'flange'
    if first in ('GRADE', 'DEGRAU'): return C('GRADE DE PISO', 'FABRICA DE GRADE DE PISO', 'AÇOS VITAL'), 'grade/degrau'
    if first in ('CHAPA', 'PLACA') and has('EXPANDIDA'): return C('CHAPA EXPANDIDA', 'FABRICA CHAPA EXPANDIDA', 'TERCEIROS'), 'chapa expandida'
    if has('PERFURADA', 'PERFURADO'): return C('PERFURADA'), 'perfurada'
    if first in ('CHAPA', 'PLACA') and has('DESGASTE', 'ANTIDESGASTE', 'ANTI DESGASTE'): return C('HARDOX', 'FABRICA USINAGEM', 'HRM'), 'chapa desgaste'
    if has('HARDOX', 'USIAR', 'USI AR'): return C('HARDOX'), 'hardox/usiar'
    if first == 'TELHA':
        # antigo: sanduíche/trapezoidal/TR40 de aço (galv/galvalume/zincado) = Fábrica de Telha; alumínio = revenda
        if not has('ALUMINIO') and (has('SANDUICHE', 'TRAPEZOIDAL', 'TR40', 'TR', 'GALVALUME', 'ZINCADO') or re.search(r'\bGALV', d)):
            return C('TELHAS', 'FABRICA DE TELHA', 'UBERABA'), 'telha (fábrica)'
        return C('TELHAS'), 'telha'
    if first in ('TUBO', 'TUBOS'): return C('TUBOS'), 'tubo'
    if first in ('CHAPA', 'CHAPAS', 'BOBINA', 'BLOCO', 'PRANCHAO', 'FITA'): return C('CHAPAS'), 'chapa'
    if first in ('PARAFUSO', 'PORCA', 'ARRUELA', 'GRAMPO', 'CHUMBADOR', 'ABRACADEIRA', 'PRISIONEIRO', 'ESTOJO', 'REBITE', 'PINO', 'CONTRAPINO', 'TIRANTE') \
       or (first == 'BARRA' and has('ROSCADA')):
        return C('FIXAÇÃO'), 'fixação'
    if first in ('PERFIL', 'CANTONEIRA', 'BARRA', 'TARUGO', 'TRILHO', 'VIGA', 'ARMADURA', 'VERGALHAO', 'FERRO', 'METALON', 'U', 'W'):
        return C('LAMINADOS'), 'laminado'
    if first in ('VALVULA', 'REGISTRO', 'PURGADOR', 'FILTRO'): return C('VALVULA'), 'válvula'
    if first in ('JUNTA', 'JUNTAS', 'GAXETA', 'ANEL') and has('JUNTA', 'ESPIRALADA', 'VEDACAO', 'GAXETA', 'O RING', 'ORING'): return C('JUNTAS'), 'junta'
    if first in ('JUNTA', 'JUNTAS'): return C('JUNTAS'), 'junta'
    if first in ('CURVA', 'COTOVELO', 'TE', 'TEE', 'REDUCAO', 'NIPLE', 'LUVA', 'PESTANA', 'UNIAO', 'MEIA', 'BUCHA', 'BUJAO', 'TAMPAO', 'CAP', 'CRUZETA', 'JOELHO', 'ANILHA', 'CONECTOR', 'ADAPTADOR', 'NIPEL', 'WELDOLET', 'SOCKOLET', 'THREADOLET', 'SWAGE', 'ESPIGAO', 'CONEXAO', 'Y'):
        return C('CONEXÕES'), 'conexão'
    if first in ('TELA', 'TELAS'): return C('TELAS'), 'tela'
    if first in ('TINTA', 'PRIMER', 'THINNER', 'VERNIZ'): return C('TINTAS'), 'tinta'
    return C('DIVERSOS'), None  # sem regra → revisar

def classificar(desc):
    c = cod_de(desc)
    if c in MAPA:
        f, l, e = MAPA[c]
        if not f:  # antigo deixou a família em branco → família pela regra (ou Caldeiraria)
            f = regra(desc)[0][0] if regra(desc)[0][1] == l else ('CALDEIRARIA' if 'CALDEIRARIA' in l else 'DIVERSOS')
        return (f, l, e), 'código já classificado no relatório anterior'
    cl, motivo = regra(desc)
    return cl, (f'regra: {motivo}' if motivo else 'SEM REGRA — conferir')

if __name__ == '__main__':
    novo = [r for r in openpyxl.load_workbook(NOVO, data_only=True).worksheets[0].iter_rows(min_row=4, values_only=True) if r[0]]
    pv = [r for r in novo if tipo_doc(r[2]) == 'Pedido de Venda']
    orig = collections.Counter()
    val = collections.Counter()
    for r in pv:
        cl, m = classificar(r[0])
        k = m.split(':')[0] if m.startswith('regra') else m
        orig[k] += 1; val[k] += float(r[1] or 0)
    for k in orig: print(k, orig[k], round(val[k], 2))
    semregra = [r for r in pv if classificar(r[0])[1].startswith('SEM')]
    c2 = collections.Counter()
    for r in semregra: c2[str(r[0]).split(' - ', 1)[-1].split()[0].upper()] += float(r[1] or 0)
    print(c2.most_common(60))
