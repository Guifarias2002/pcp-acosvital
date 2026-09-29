# Organiza os relatórios mensais em  <PASTA>/2026/<n> - <Mês>/  com Excel + PDF
# com as MESMAS informações (COM Flanges). Excel no layout do modelo CONTROLE GERAL
# (aba "Faturamento por Produto") + aba "Resumo" com os mesmos quadros do PDF.
# Agosto  = CONTROLE GERAL original (relatório do dia 19), com Flanges.
# Setembro = abas "Setembro sem Flanges" + "Flanges (retirados)" do classificado (Omie + Caldeiraria do sistema).
import collections, os, shutil, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from classif import norm
from pdf_modelo import PASTA, gerar

H = Font(bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1A3A5C')
VD = PatternFill('solid', fgColor='EAF6EF'); LR = PatternFill('solid', fgColor='FDF1E8')
BRL = '"R$" #,##0.00'; PCT = '0.0%'
tit = lambda s: str(s).title().replace('Ç', 'ç')


def excel(reg, saida, titulo, extras=None):
    """reg = [(desc, valor, pedido, nf, familia, local, empresa)]"""
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Faturamento por Produto'
    ws.append([titulo]); ws['A1'].font = Font(bold=True, size=14, color='1A3A5C')
    ws.append(['Descrição do Produto (completa)', 'Total da Nota Fiscal', 'Pedido', 'Nota Fiscal', 'Familia ', 'Local de produção ', 'Empresa de fabricação'])
    for c in ws[2]: c.font = H; c.fill = HF
    n = len(reg)
    ws.append([None, f'=SUM(B4:B{n + 3})']); ws['B3'].font = Font(bold=True); ws['B3'].number_format = BRL
    for r in sorted(reg, key=lambda r: str(r[0])):
        ws.append(list(r)); ws.cell(ws.max_row, 2).number_format = '#,##0.00'
    ws.auto_filter.ref = f'A2:G{n + 3}'; ws.freeze_panes = 'A4'
    for i, w in enumerate([80, 18, 26, 12, 18, 28, 20], 1): ws.column_dimensions[get_column_letter(i)].width = w

    # ── Resumo = mesmos quadros do PDF ───────────────────────────────────────
    E = collections.Counter(); L = collections.Counter(); F = collections.Counter(); LF = collections.defaultdict(collections.Counter)
    for _, v, _, _, f, l, e in reg:
        E[e] += v; L[l] += v; F[f] += v; LF[l][f] += v
    tot = sum(r[1] for r in reg); fab = tot - L['REVENDA']
    s = wb.create_sheet('Resumo', 0)
    s.append([titulo]); s['A1'].font = Font(bold=True, size=14, color='1A3A5C')
    s.append([]); s.append(['FATURAMENTO TOTAL', tot]); s.append(['FABRICADO (produção própria)', fab, fab / tot if tot else 0])
    s.append(['REVENDA', L['REVENDA'], L['REVENDA'] / tot if tot else 0])
    for i in (3, 4, 5):
        s.cell(i, 1).font = Font(bold=True); s.cell(i, 2).number_format = BRL; s.cell(i, 3).number_format = PCT

    def quadro(cab, linhas, total=None, pinta=()):
        s.append([]); s.append(cab)
        for c in s[s.max_row]: c.font = H; c.fill = HF
        for k, v, p in linhas:
            s.append([tit(k), v, p]); s.cell(s.max_row, 2).number_format = BRL; s.cell(s.max_row, 3).number_format = PCT
            if k in dict(pinta):
                for c in s[s.max_row]: c.fill = dict(pinta)[k]; c.font = Font(bold=True)
        if total:
            s.append(total); s.cell(s.max_row, 2).number_format = BRL; s.cell(s.max_row, 3).number_format = PCT
            for c in s[s.max_row]: c.font = H; c.fill = HF

    quadro(['Empresa de Fabricação', 'Valor (NF)', '% do Total'], [(k, v, v / tot) for k, v in E.most_common()], ['TOTAL GERAL', tot, 1])
    Lfab = collections.Counter({k: v for k, v in L.items() if k != 'REVENDA'})
    quadro(['Local de Produção', 'Valor (NF)', '% do Fabricado'],
           [(k, v, v / fab) for k, v in Lfab.most_common()] + [('TOTAL FABRICADO', fab, 1), ('Revenda (não fabricado)', L['REVENDA'], None)],
           ['TOTAL GERAL FATURADO', tot, None], pinta=[('TOTAL FABRICADO', VD), ('Revenda (não fabricado)', LR)])
    quadro(['Família', 'Valor (NF)', '% do Total'], [(k, v, v / tot) for k, v in F.most_common()], ['TOTAL', tot, 1])
    s.append([]); s.append(['Detalhe — Família dentro de cada Fábrica']); s.cell(s.max_row, 1).font = Font(bold=True, size=12, color='1A3A5C')
    for loc, v in Lfab.most_common():
        quadro([tit(loc), v, '% da fábrica'], [(k, x, x / v if v else 0) for k, x in LF[loc].most_common()])
        s.cell(s.max_row - len(LF[loc]), 2).number_format = BRL
    s.append([]); s.append(['“Fabricado” = todos os Locais de Produção exceto Revenda. Valores por Total da Nota Fiscal.'])
    s.cell(s.max_row, 1).alignment = Alignment(wrap_text=False); s.cell(s.max_row, 1).font = Font(italic=True, color='64748B')
    for col, w in zip('ABC', [44, 20, 14]): s.column_dimensions[col].width = w

    for nome, fonte in (extras or []):  # abas de conferência copiadas como estão
        d = wb.create_sheet(nome[:31])
        for row in fonte.iter_rows(values_only=True): d.append(list(row))
        for c in d[1]: c.font = H; c.fill = HF
        for col, dim in fonte.column_dimensions.items(): d.column_dimensions[col].width = dim.width
    wb.save(saida)
    return tot, fab


def pasta(n, mes):
    # FATURAMENTO_SAIDA = outra pasta raiz (ex.: quando o C: está cheio, grava na do sistema no Z:)
    p = os.path.join(os.environ.get('FATURAMENTO_SAIDA', PASTA), '2026', f'{n} - {mes}'); os.makedirs(p, exist_ok=True); return p


if __name__ == '__main__':
    # ── Agosto ───────────────────────────────────────────────────────────────────
    ORIG = PASTA + 'CONTROLE GERAL DE FATURAMENTO AÇOS VITAL.xlsx'
    ag = [(str(r[0]), float(r[1] or 0), r[2], r[3], norm(r[4]) or '(SEM FAMÍLIA)', norm(r[5]), norm(r[6]))
          for r in openpyxl.load_workbook(ORIG, data_only=True).worksheets[0].iter_rows(min_row=4, values_only=True) if r[0]]
    pa = pasta(8, 'Agosto')
    print('agosto', excel(ag, os.path.join(pa, 'Faturamento - Agosto 2026.xlsx'), 'Faturamento por Produto — Agosto 2026'))
    print(gerar([(r[1], r[4], r[5], r[6], str(r[3])) for r in ag], os.path.join(pa, 'Faturamento - Agosto 2026.pdf'),
                'Agosto 2026 · Por Empresa de Fabricação e Local de Produção · valores por Total da Nota Fiscal',
                'Fonte: CONTROLE GERAL DE FATURAMENTO AÇOS VITAL.xlsx (relatório de agosto, apresentado no dia 19). Inclui Flanges.'))

    # ── Setembro ─────────────────────────────────────────────────────────────────
    wbS = openpyxl.load_workbook(PASTA + 'FATURAMENTO SETEMBRO (sem Flanges) - CLASSIFICADO.xlsx', data_only=True)
    se = [(str(r[0]), float(r[1] or 0), r[2], r[3], r[4], r[5], r[6])
          for aba in ('Setembro sem Flanges', 'Flanges (retirados)')
          for r in wbS[aba].iter_rows(min_row=2, values_only=True) if r[0] and r[0] != 'TOTAL']
    ps = pasta(9, 'Setembro')
    print('setembro', excel(se, os.path.join(ps, 'Faturamento - Setembro 2026.xlsx'), 'Faturamento por Produto — Setembro 2026',
                            extras=[(n, wbS[n]) for n in wbS.sheetnames if n not in ('Setembro sem Flanges', 'Flanges (retirados)')]))
    print(gerar([(r[1], r[4], r[5], r[6], str(r[3])) for r in se], os.path.join(ps, 'Faturamento - Setembro 2026.pdf'),
                'Setembro 2026 · planilha Omie + Caldeiraria do sistema · Por Empresa de Fabricação e Local de Produção · valores por Total da Nota Fiscal',
                'Fonte: export Omie de setembro (só Pedido de Venda), mesmos critérios do CONTROLE GERAL, com Flanges. Caldeiraria (Planejamento) = sistema PCP Caldeiraria '
                '(mês de chegada na Caldeiraria); linhas da planilha de pedidos que já estão no sistema foram retiradas para não contar duas vezes.'))
